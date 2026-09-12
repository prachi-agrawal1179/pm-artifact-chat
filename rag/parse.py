from __future__ import annotations

from pathlib import Path
from urllib.parse import urlparse

import httpx

SUPPORTED_SUFFIXES = {
    ".md",
    ".txt",
    ".markdown",
    ".pdf",
    ".html",
    ".htm",
    ".docx",
    ".xlsx",
    ".xlsm",
    ".xls",
    ".csv",
}


def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(path)
    if suffix == ".xls":
        return _parse_xls(path)
    if suffix == ".csv":
        return _parse_csv(path)
    return path.read_text(encoding="utf-8", errors="replace")


def parse_url(url: str, timeout: float = 20.0) -> tuple[str, str]:
    """Fetch a public page and return (source_label, text)."""
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise ValueError("Only http(s) URLs are supported.")
    response = httpx.get(
        url,
        timeout=timeout,
        follow_redirects=True,
        headers={"User-Agent": "pm-rag-chatbot/1.0 (local demo)"},
    )
    response.raise_for_status()
    content_type = response.headers.get("content-type", "")
    path_lower = parsed.path.lower()
    if "pdf" in content_type or path_lower.endswith(".pdf"):
        from io import BytesIO

        from pypdf import PdfReader

        reader = PdfReader(BytesIO(response.content))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif path_lower.endswith(".docx") or "wordprocessingml" in content_type:
        from tempfile import NamedTemporaryFile

        with NamedTemporaryFile(suffix=".docx") as tmp:
            tmp.write(response.content)
            tmp.flush()
            text = _parse_docx(Path(tmp.name))
    elif path_lower.endswith((".xlsx", ".xlsm")) or "spreadsheetml" in content_type:
        from tempfile import NamedTemporaryFile

        suffix = ".xlsm" if path_lower.endswith(".xlsm") else ".xlsx"
        with NamedTemporaryFile(suffix=suffix) as tmp:
            tmp.write(response.content)
            tmp.flush()
            text = _parse_xlsx(Path(tmp.name))
    else:
        text = _html_to_text(response.text)
    label = parsed.netloc + parsed.path
    return label, text.strip()


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(path: Path) -> str:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    parts: list[str] = []
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style else ""
            if style.startswith("Heading"):
                digits = "".join(ch for ch in style if ch.isdigit())
                level = int(digits) if digits else 1
                parts.append("#" * min(level, 6) + " " + text)
            else:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _stringify_cell(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _sheet_to_text(title: str, rows: list[list[str]]) -> str:
    nonempty = [row for row in rows if any(cell.strip() for cell in row)]
    if not nonempty:
        return ""
    lines = [f"# {title}", ""]
    lines.extend(" | ".join(row) for row in nonempty)
    return "\n".join(lines)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([_stringify_cell(cell) for cell in row])
            text = _sheet_to_text(sheet.title, rows)
            if text:
                sheets.append(text)
    finally:
        workbook.close()
    return "\n\n".join(sheets)


def _parse_xls(path: Path) -> str:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            rows.append([_stringify_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)])
        text = _sheet_to_text(sheet.name, rows)
        if text:
            sheets.append(text)
    return "\n\n".join(sheets)


def _parse_csv(path: Path) -> str:
    import csv

    with path.open(encoding="utf-8", errors="replace", newline="") as handle:
        rows = [[_stringify_cell(cell) for cell in row] for row in csv.reader(handle)]
    return _sheet_to_text(path.stem, rows)


def _html_to_text(html: str) -> str:
    import re

    html = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>", "\n", html)
    html = re.sub(r"(?i)</p>", "\n\n", html)
    html = re.sub(r"(?i)</h[1-6]>", "\n\n", html)
    html = re.sub(r"(?s)<[^>]+>", " ", html)
    html = re.sub(r"[ \t]+", " ", html)
    html = re.sub(r"\n{3,}", "\n\n", html)
    return html


def iter_doc_paths(root: Path) -> list[Path]:
    paths: list[Path] = []
    if not root.exists():
        return paths
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            paths.append(path)
    return paths
